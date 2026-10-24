#!/usr/bin/env python3
"""Turn the auto-trader's journal.json into trade-journal.xlsx and trade-journal.pdf (A4 landscape).

Excel sheets
  Summary      key numbers as live formulas over the other sheets
  Orders       one row per order, every detail recorded when it was placed (columns A to AV)
  Round trips  each buy lot matched first-in-first-out to its sells, with P/L as formulas
  Rules        the auto-trader rules and the four checks

Times are written three ways: New York (exchange time), Berlin and UTC.

Usage: python3 journal_report.py journal.json trades.json OUT_DIR
Needs: openpyxl, reportlab
"""
import datetime as dt
import json
import os
import sys
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NY, BERLIN, UTC = ZoneInfo("America/New_York"), ZoneInfo("Europe/Berlin"), dt.timezone.utc
TEAL, INK, MUTED, LINE, SOFT = "0D6E8A", "132130", "627284", "D7DEE5", "F3F6F8"
UP, DOWN = "12805C", "C2362F"
F = "Arial"


def parse(x):
    if not x:
        return None
    try:
        t = dt.datetime.fromisoformat(str(x).replace("Z", "+00:00"))
        return t if t.tzinfo else t.replace(tzinfo=UTC)
    except ValueError:
        return None


def local(t, tz):
    return t.astimezone(tz).replace(tzinfo=None, microsecond=0) if t else None


def when(e):
    return parse(e.get("filled_at_utc")) or parse(e.get("submitted_at_utc")) or parse(e.get("decided_at_utc"))


def round_trips(entries):
    """FIFO-match filled buys to filled sells per symbol."""
    lots, rows = {}, []
    for e in sorted((e for e in entries if e.get("filled_qty") and e.get("filled_avg_price")), key=lambda e: when(e) or dt.datetime.min.replace(tzinfo=UTC)):
        sym, qty, px = e["symbol"], float(e["filled_qty"]), float(e["filled_avg_price"])
        if e["side"] == "buy":
            lots.setdefault(sym, []).append({"symbol": sym, "company": e.get("company"), "bought": when(e), "buy_px": px, "qty": qty})
            continue
        q = lots.get(sym, [])
        while qty > 1e-12 and q:
            lot = q[0]
            take = min(qty, lot["qty"])
            rows.append(dict(lot, qty=take, sold=when(e), sell_px=px, exit=e.get("rule") or e.get("reason"), status="Closed"))
            lot["qty"] -= take
            qty -= take
            if lot["qty"] <= 1e-12:
                q.pop(0)
    for q in lots.values():
        for lot in q:
            rows.append(dict(lot, sold=None, sell_px=None, exit=None, status="Open"))
    rows.sort(key=lambda r: (r["bought"] or dt.datetime.min.replace(tzinfo=UTC), r["symbol"]))
    return rows


# ----------------------------------------------------------------------------- Excel
ORDER_COLS = [
    # (header, key or callable, width, number format)
    ("No", "journal_no", 6, "0"),
    ("Time Berlin", lambda e: local(when(e), BERLIN), 19, "yyyy-mm-dd hh:mm:ss"),
    ("Time New York", lambda e: local(when(e), NY), 19, "yyyy-mm-dd hh:mm:ss"),
    ("Time UTC", lambda e: local(when(e), UTC), 19, "yyyy-mm-dd hh:mm:ss"),
    ("Side", lambda e: (e.get("side") or "").upper(), 7, None),
    ("Symbol", "symbol", 9, None),
    ("Company", "company", 28, None),
    ("Exchange", "exchange", 10, None),
    ("Status", "status", 11, None),
    ("Qty filled", "filled_qty", 12, "0.000000"),
    ("Avg fill price $", "filled_avg_price", 13, "$#,##0.00"),
    ("Fill value $", None, 13, "$#,##0.00"),                      # formula J*K
    ("Requested $", "notional_requested", 12, "$#,##0.00"),
    ("Requested qty", "qty_requested", 12, "0.000000"),
    ("Rule", "rule", 18, None),
    ("Reason", "reason", 40, None),
    ("Greens (of 4)", "greens", 9, "0"),
    ("Agent rank", "agent_rank", 8, "0"),
    ("1 Data real", "check1_data_status", 10, None),
    ("Source spread %", lambda e: (e.get("check1_spread_pct") or 0) / 100 if e.get("check1_spread_pct") is not None else None, 10, "0.00%"),
    ("2 Organic", "check2_organic_status", 10, None),
    ("Organic score", "check2_score", 9, "0"),
    ("Organic flags", "check2_flags", 40, None),
    ("3 Prediction", "check3_prediction_status", 11, None),
    ("Call", "check3_call", 7, None),
    ("Target $", "check3_target", 11, "$#,##0.00"),
    ("Range low $", "check3_low", 11, "$#,##0.00"),
    ("Range high $", "check3_high", 11, "$#,##0.00"),
    ("Call due", lambda e: dt.date.fromisoformat(e["check3_due"]) if e.get("check3_due") else None, 11, "yyyy-mm-dd"),
    ("Confidence", "check3_confidence", 10, "0%"),
    ("4 News", "check4_news_status", 9, None),
    ("Headlines (72h)", "check4_headlines", 9, "0"),
    ("Top headline", "check4_top_headline", 44, None),
    ("Price at decision $", "price_at_decision", 13, "$#,##0.00"),
    ("Move today %", lambda e: e["move_today_pct"] / 100 if e.get("move_today_pct") is not None else None, 10, "0.00%"),
    ("Account equity before $", "account_equity_before", 15, "$#,##0.00"),
    ("Cash before $", "account_cash_before", 13, "$#,##0.00"),
    ("Slot %", lambda e: e["slot_pct"] / 100 if e.get("slot_pct") is not None else None, 8, "0.00%"),
    ("Position qty before", "position_qty_before", 12, "0.000000"),
    ("Avg entry before $", "position_avg_entry", 13, "$#,##0.00"),
    ("P/L % before", lambda e: e["position_plpc_before"] / 100 if e.get("position_plpc_before") is not None else None, 10, "0.00%"),
    ("Account", "mode", 8, None),
    ("Order ID", "order_id", 38, None),
    ("Client order ID", "client_order_id", 40, None),
    ("Decided at (UTC)", lambda e: local(parse(e.get("decided_at_utc")), UTC), 19, "yyyy-mm-dd hh:mm:ss"),
    ("Filled at (UTC)", lambda e: local(parse(e.get("filled_at_utc")), UTC), 19, "yyyy-mm-dd hh:mm:ss"),
    ("Agent run (UTC)", lambda e: local(parse(e.get("agent_run_at")), UTC), 19, "yyyy-mm-dd hh:mm:ss"),
    ("Error", "error", 40, None),
]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def add_table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def build_xlsx(entries, trades, path, generated):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    base = Font(name=F, size=10)

    # ---------------- Orders
    wo = wb.active
    wo.title = "Orders"
    headers = [c[0] for c in ORDER_COLS]
    wo.append(headers)
    style_header(wo, 1, len(headers))
    fv = headers.index("Fill value $") + 1
    qcol, pcol = get_column_letter(headers.index("Qty filled") + 1), get_column_letter(headers.index("Avg fill price $") + 1)
    rows = sorted(entries, key=lambda e: e.get("journal_no") or 0)
    for r, e in enumerate(rows, start=2):
        for c, (h, key, w, fmt) in enumerate(ORDER_COLS, start=1):
            if h == "Fill value $":
                v = f'=IF(AND(ISNUMBER({qcol}{r}),ISNUMBER({pcol}{r})),{qcol}{r}*{pcol}{r},"")'
            else:
                v = key(e) if callable(key) else e.get(key)
            cell = wo.cell(row=r, column=c, value=v)
            cell.font = base
            if fmt:
                cell.number_format = fmt
            if h == "Side":
                cell.font = Font(name=F, size=10, bold=True, color=UP if v == "BUY" else DOWN)
            if h in ("Reason", "Organic flags", "Top headline", "Company"):
                cell.alignment = Alignment(wrap_text=False)
    last = max(2, len(rows) + 1)
    for c, (h, key, w, fmt) in enumerate(ORDER_COLS, start=1):
        wo.column_dimensions[get_column_letter(c)].width = w
    wo.freeze_panes = "G2"
    if rows:
        add_table(wo, "Orders", f"A1:{get_column_letter(len(headers))}{last}")
    else:
        wo.cell(row=2, column=1, value="No orders yet. Rows appear here as soon as the auto-trader places one.").font = Font(name=F, italic=True, color=MUTED)
    side_col = get_column_letter(headers.index("Side") + 1)
    status_col = get_column_letter(headers.index("Status") + 1)
    fv_col = get_column_letter(fv)

    # ---------------- Round trips
    wr = wb.create_sheet("Round trips")
    rh = ["Symbol", "Company", "Bought (Berlin)", "Bought (New York)", "Buy price $", "Qty", "Cost $",
          "Sold (Berlin)", "Sold (New York)", "Sell price $", "Proceeds $", "P/L $", "P/L %", "Days held", "Exit rule", "Status"]
    wr.append(rh)
    style_header(wr, 1, len(rh))
    trips = round_trips(entries)
    for r, t in enumerate(trips, start=2):
        vals = [t["symbol"], t.get("company"), local(t["bought"], BERLIN), local(t["bought"], NY), t["buy_px"], t["qty"],
                f"=E{r}*F{r}", local(t["sold"], BERLIN), local(t["sold"], NY), t["sell_px"],
                f'=IF(ISNUMBER(J{r}),J{r}*F{r},"")', f'=IF(ISNUMBER(J{r}),K{r}-G{r},"")', f'=IF(ISNUMBER(J{r}),L{r}/G{r},"")',
                f'=IF(ISNUMBER(H{r}),H{r}-C{r},"")', t.get("exit"), t["status"]]
        fmts = [None, None, "yyyy-mm-dd hh:mm", "yyyy-mm-dd hh:mm", "$#,##0.00", "0.000000", "$#,##0.00",
                "yyyy-mm-dd hh:mm", "yyyy-mm-dd hh:mm", "$#,##0.00", "$#,##0.00", "$#,##0.00;[Red]-$#,##0.00", "0.00%;[Red]-0.00%", "0.0", None, None]
        for c, (v, fmt) in enumerate(zip(vals, fmts), start=1):
            cell = wr.cell(row=r, column=c, value=v)
            cell.font = base
            if fmt:
                cell.number_format = fmt
    rlast = max(2, len(trips) + 1)
    for c, w in enumerate([9, 28, 17, 17, 12, 12, 12, 17, 17, 12, 12, 12, 10, 10, 20, 9], start=1):
        wr.column_dimensions[get_column_letter(c)].width = w
    wr.freeze_panes = "B2"
    if trips:
        add_table(wr, "RoundTrips", f"A1:P{rlast}")
        green = Font(name=F, size=10, color=UP)
        wr.conditional_formatting.add(f"L2:M{rlast}", CellIsRule(operator="greaterThan", formula=["0"], font=green))
    else:
        wr.cell(row=2, column=1, value="No filled buys yet.").font = Font(name=F, italic=True, color=MUTED)

    # ---------------- Summary
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "Trade Journal · Trade Desk"
    ws["A1"].font = Font(name=F, size=16, bold=True, color=INK)
    ws["A2"] = "Generated (Berlin)"
    ws["B2"] = local(generated, BERLIN)
    ws["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["A3"] = "Generated (New York)"
    ws["B3"] = local(generated, NY)
    ws["B3"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["A4"] = "Account"
    ws["B4"] = "Paper (practice money)" if trades.get("paper", True) else "LIVE (real money)"
    ws["A5"] = "Auto-trader switch"
    ws["B5"] = (trades.get("mode") or "off").upper()
    O, R = f"Orders!${side_col}$2:${side_col}${last}", "'Round trips'"
    st = f"Orders!${status_col}$2:${status_col}${last}"
    val = f"Orders!${fv_col}$2:${fv_col}${last}"
    kpis = [
        ("Orders placed", f'=COUNTIF({O},"BUY")+COUNTIF({O},"SELL")', "0"),
        ("Buy orders", f'=COUNTIF({O},"BUY")', "0"),
        ("Sell orders", f'=COUNTIF({O},"SELL")', "0"),
        ("Filled orders", f'=COUNTIF({st},"filled")', "0"),
        ("Total bought $", f'=SUMIFS({val},{O},"BUY")', "$#,##0.00"),
        ("Total sold $", f'=SUMIFS({val},{O},"SELL")', "$#,##0.00"),
        ("Closed round trips", f'=COUNTIF({R}!$P$2:$P${rlast},"Closed")', "0"),
        ("Open lots", f'=COUNTIF({R}!$P$2:$P${rlast},"Open")', "0"),
        ("Realized P/L $", f"=SUM({R}!$L$2:$L${rlast})", "$#,##0.00;[Red]-$#,##0.00"),
        ("Winning round trips", f'=COUNTIFS({R}!$P$2:$P${rlast},"Closed",{R}!$L$2:$L${rlast},">0")', "0"),
        ("Win rate", "=IFERROR(B17/B14,0)", "0.0%"),
    ]
    ws["A7"], ws["B7"] = "Measure", "Value"
    style_header(ws, 7, 2)
    for i, (k, f, fmt) in enumerate(kpis, start=8):
        ws.cell(row=i, column=1, value=k).font = base
        c = ws.cell(row=i, column=2, value=f)
        c.font = Font(name=F, size=10, bold=True)
        c.number_format = fmt
    ws["A20"] = "Times: New York is the exchange clock; Berlin is your local time. Round trips match buys to sells first-in-first-out."
    ws["A21"] = "Source: Alpaca order data recorded by the auto-trader at the moment each order was placed."
    for a in ("A20", "A21"):
        ws[a].font = Font(name=F, size=9, italic=True, color=MUTED)
    for a in ("A2", "A3", "A4", "A5"):
        ws[a].font = Font(name=F, size=10, color=MUTED)
    for a in ("B2", "B3", "B4", "B5"):
        ws[a].font = Font(name=F, size=10, bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 26

    # ---------------- Rules
    wu = wb.create_sheet("Rules")
    R2 = trades.get("rules") or {}
    rules = [
        ("Buy", "Only stocks rated 4 of 4 green that are not already held; funds/ETFs skipped" + (" (allowed)" if R2.get("allow_etf") else "")),
        ("Slot size", f"Account split into {R2.get('slots', 12)} equal slots of {R2.get('slot_pct', 8.33)}% each; cash only, never margin"),
        ("No buys before close", f"Last {R2.get('no_buy_minutes_before_close', 30):g} minutes of the session: sells only"),
        ("Daily loss guard", f"No new buys on a day the account is down {R2.get('daily_loss_pct', 5):g}% or more"),
        ("Exit: 3 green", "Sell half, once (not again within 3 days)"),
        ("Exit: 2 or fewer green", "Sell everything"),
        ("Stop-loss", f"Down {R2.get('stop_loss_pct', 8):g}% from entry: sell everything"),
        ("Take profit", f"Up {R2.get('take_profit_pct', 20):g}% from entry: sell everything"),
        ("Scope", "Only positions the auto-trader bought (client order ID starts with tdagent-)"),
        ("Check 1 · Data is real", "Price agrees across trade, quote, minute, hourly and daily bars; history has no broken bars or holes"),
        ("Check 2 · Organic", "No volume pump without news, no one-day jump doing the work, no gaps, no pump-and-fade, not thin"),
        ("Check 3 · Prediction", "5-trading-day momentum call with a volatility range; green only when it leans up"),
        ("Check 4 · News", "Headlines in the last 72 hours scored for high-impact words against today's move"),
    ]
    wu.append(["Rule", "What it does"])
    style_header(wu, 1, 2)
    for k, v in rules:
        wu.append([k, v])
    for row in wu.iter_rows(min_row=2):
        for c in row:
            c.font = base
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wu.column_dimensions["A"].width = 26
    wu.column_dimensions["B"].width = 100

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = sh.title not in ("Summary", "Rules")
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return trips


# ----------------------------------------------------------------------------- PDF
def build_pdf(entries, trips, trades, path, generated):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as T, TableStyle

    c = lambda h: colors.HexColor("#" + h)
    small = ParagraphStyle("s", fontName="Helvetica", fontSize=7.4, leading=9, textColor=c(INK))
    head = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=7.4, leading=9, textColor=colors.white)
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=c(INK))
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=c(INK), spaceBefore=8, spaceAfter=4)
    meta = ParagraphStyle("m", fontName="Helvetica", fontSize=8.5, leading=11, textColor=c(MUTED))
    kv = ParagraphStyle("kv", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=c(INK))
    kk = ParagraphStyle("kk", fontName="Helvetica", fontSize=7, leading=9, textColor=c(MUTED))

    money = lambda v: "" if v is None else f"${v:,.2f}"
    qtyf = lambda v: "" if v is None else (f"{v:,.6f}".rstrip("0").rstrip("."))
    tfmt = lambda t, tz: t.astimezone(tz).strftime("%Y-%m-%d %H:%M") if t else ""
    P = lambda s, st=small: Paragraph(str(s) if s is not None else "", st)

    filled = [e for e in entries if e.get("filled_qty") and e.get("filled_avg_price")]
    bought = sum(e["filled_qty"] * e["filled_avg_price"] for e in filled if e["side"] == "buy")
    sold = sum(e["filled_qty"] * e["filled_avg_price"] for e in filled if e["side"] == "sell")
    closed = [t for t in trips if t["status"] == "Closed"]
    pl = sum((t["sell_px"] - t["buy_px"]) * t["qty"] for t in closed)
    wins = sum(1 for t in closed if t["sell_px"] > t["buy_px"])

    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title="Trade Journal", author="Trade Desk auto-trader")
    W = landscape(A4)[0] - 24 * mm
    story = [P("Trade Journal", h1),
             P(f"Trade Desk auto-trader · {'Paper account (practice money)' if trades.get('paper', True) else 'LIVE account (real money)'} · "
               f"switch {str(trades.get('mode') or 'off').upper()} · generated {tfmt(generated, BERLIN)} Berlin / {tfmt(generated, NY)} New York", meta),
             Spacer(1, 6)]
    kp = [("Orders", f"{len(entries)}"), ("Filled", f"{len(filled)}"), ("Bought", money(bought)), ("Sold", money(sold)),
          ("Closed trips", f"{len(closed)}"), ("Realized P/L", money(pl)), ("Win rate", f"{(wins / len(closed) * 100):.0f}%" if closed else "—"),
          ("Open lots", f"{sum(1 for t in trips if t['status'] == 'Open')}")]
    kt = T([[P(v, kv) for _, v in kp], [P(k, kk) for k, _ in kp]], colWidths=[W / len(kp)] * len(kp))
    kt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c(SOFT)), ("BOX", (0, 0), (-1, -1), 0.5, c(LINE)),
                            ("LINEAFTER", (0, 0), (-2, -1), 0.5, c(LINE)), ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story += [kt, Spacer(1, 6)]

    def table(header, rows, widths, numeric_cols=()):
        data = [[P(h, head_num if i in numeric_cols else head) for i, h in enumerate(header)]] + rows
        t = T(data, colWidths=[W * w for w in widths], repeatRows=1)
        stl = [("BACKGROUND", (0, 0), (-1, 0), c(TEAL)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
               ("LINEBELOW", (0, 0), (-1, -1), 0.25, c(LINE)), ("TOPPADDING", (0, 0), (-1, -1), 3),
               ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        for i in range(1, len(data)):
            if i % 2 == 0:
                stl.append(("BACKGROUND", (0, i), (-1, i), c(SOFT)))
        t.setStyle(TableStyle(stl))
        return t

    num_style = ParagraphStyle("n", parent=small, alignment=2)
    head_num = ParagraphStyle("hn", parent=head, alignment=2)
    CODE = {"ok": "P", "watch": "W", "alert": "A", "skip": "N"}
    story.append(P("Orders", h2))
    if entries:
        rows = []
        for e in sorted(entries, key=lambda e: e.get("journal_no") or 0):
            t = when(e)
            side = (e.get("side") or "").upper()
            sc = UP if side == "BUY" else DOWN
            checks = " ".join(f"{i}{CODE.get(e.get(k), '–')}" for i, k in enumerate(
                ("check1_data_status", "check2_organic_status", "check3_prediction_status", "check4_news_status"), 1))
            rows.append([P(e.get("journal_no")), P(tfmt(t, BERLIN)), P(tfmt(t, NY)),
                         P(f'<font color="#{sc}"><b>{side}</b></font>'), P(f"<b>{e.get('symbol')}</b>"), P(e.get("company") or ""),
                         P(qtyf(e.get("filled_qty") or e.get("qty_requested")), num_style),
                         P(money(e.get("filled_avg_price")), num_style),
                         P(money(e["filled_qty"] * e["filled_avg_price"] if e.get("filled_qty") and e.get("filled_avg_price") else e.get("notional_requested")), num_style),
                         P(e.get("status") or ""), P(f"{e.get('greens', '')}/4 · {checks}"), P(e.get("reason") or "")])
        story.append(table(["#", "Time Berlin", "Time New York", "Side", "Symbol", "Company", "Qty", "Avg price", "Value",
                            "Status", "Checks at decision", "Reason"], rows,
                           [0.03, 0.085, 0.085, 0.045, 0.06, 0.15, 0.07, 0.07, 0.075, 0.065, 0.11, 0.155], numeric_cols=(6, 7, 8)))
        story.append(Spacer(1, 3))
        story.append(P("Checks at decision: 1 data is real · 2 organic trend · 3 prediction · 4 news. P pass, W watch, A alert, N not read.", meta))
    else:
        story.append(P("No orders yet. They appear here as soon as the auto-trader places one.", meta))

    story.append(P("Round trips (first-in, first-out)", h2))
    if trips:
        rows = []
        for t in trips:
            plv = (t["sell_px"] - t["buy_px"]) * t["qty"] if t["sell_px"] is not None else None
            plp = (t["sell_px"] / t["buy_px"] - 1) * 100 if t["sell_px"] is not None and t["buy_px"] else None
            col = UP if (plv or 0) >= 0 else DOWN
            days = (t["sold"] - t["bought"]).total_seconds() / 86400 if t["sold"] and t["bought"] else None
            rows.append([P(f"<b>{t['symbol']}</b>"), P(tfmt(t["bought"], BERLIN)), P(money(t["buy_px"]), num_style), P(qtyf(t["qty"]), num_style),
                         P(money(t["buy_px"] * t["qty"]), num_style), P(tfmt(t["sold"], BERLIN)), P(money(t["sell_px"]), num_style),
                         P(money(t["sell_px"] * t["qty"]) if t["sell_px"] is not None else "", num_style),
                         P(f'<font color="#{col}">{money(plv)}</font>' if plv is not None else "", num_style),
                         P(f'<font color="#{col}">{plp:+.2f}%</font>' if plp is not None else "", num_style),
                         P(f"{days:.1f}" if days is not None else "", num_style), P(t.get("exit") or ""), P(t["status"])])
        story.append(table(["Symbol", "Bought (Berlin)", "Buy price", "Qty", "Cost", "Sold (Berlin)", "Sell price", "Proceeds",
                            "P/L", "P/L %", "Days", "Exit rule", "Status"], rows,
                           [0.065, 0.095, 0.075, 0.08, 0.08, 0.095, 0.075, 0.08, 0.075, 0.06, 0.045, 0.115, 0.06], numeric_cols=(2, 3, 4, 6, 7, 8, 9, 10)))
    else:
        story.append(P("No filled buys yet.", meta))

    def footer(cv, d):
        cv.saveState()
        cv.setFont("Helvetica", 7.5)
        cv.setFillColor(c(MUTED))
        cv.drawString(12 * mm, 7 * mm, f"Trade Desk · Trade Journal · generated {tfmt(generated, BERLIN)} Berlin · not financial advice")
        cv.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Page {d.page}")
        cv.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main():
    journal_path, trades_path, out_dir = sys.argv[1:4]
    try:
        entries = json.load(open(journal_path)).get("entries", [])
    except Exception:
        entries = []
    try:
        trades = json.load(open(trades_path))
    except Exception:
        trades = {}
    generated = dt.datetime.now(UTC)
    os.makedirs(out_dir, exist_ok=True)
    trips = build_xlsx(entries, trades, os.path.join(out_dir, "trade-journal.xlsx"), generated)
    build_pdf(entries, trips, trades, os.path.join(out_dir, "trade-journal.pdf"), generated)
    print(f"journal: {len(entries)} orders, {len(trips)} lots -> trade-journal.xlsx, trade-journal.pdf")


if __name__ == "__main__":
    main()
